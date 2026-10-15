"""Strict Excel invoice intake: explicit columns, no guessed sales, bounded resource use."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import cast

from openpyxl import load_workbook

from office_assistant.application.sales_ledger import SaleDraft, SaleValidationError

MAX_XLSX_BYTES = 2 * 1024 * 1024
MAX_INVOICES = 100
_FIELDS: dict[str, frozenset[str]] = {
    "customer": frozenset({"müştəri", "şirkət", "customer", "company"}),
    "date": frozenset({"tarix", "date", "satış tarixi", "invoice date"}),
    "invoice": frozenset({"faktura", "faktura nömrəsi", "invoice", "invoice number"}),
    "amount": frozenset({"məbləğ", "satış məbləği", "amount", "total"}),
    "currency": frozenset({"valyuta", "currency"}),
}


def _header(value: object) -> str:
    return " ".join(str(value).casefold().strip().split())


def _date(value: object, row: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError as error:
            raise SaleValidationError(f"{row}-ci sətr: tarixi YYYY-MM-DD yazın.") from error
    raise SaleValidationError(f"{row}-ci sətr: tarix tanınmadı.")


def parse_sales_xlsx(
    content: bytes, *, today: date | None = None
) -> tuple[SaleDraft, ...]:
    """Return a preview only. Never persist a workbook without explicit human approval.

    Exactly one completed invoice per row. Mixed currencies are allowed but never
    summed together. Formula amounts are rejected instead of trusting cached cells.
    """

    if not content or len(content) > MAX_XLSX_BYTES:
        raise SaleValidationError("Excel faylı boşdur və ya 2 MiB həddini aşır.")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (OSError, ValueError, KeyError) as error:
        raise SaleValidationError("XLSX faylı oxuna bilmir.") from error
    try:
        visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        if len(visible) != 1:
            raise SaleValidationError("İdxal üçün yalnız bir görünən Excel vərəqi olmalıdır.")
        sheet = visible[0]
        rows = sheet.iter_rows(values_only=True)
        headings = next(rows, None)
        if headings is None:
            raise SaleValidationError("Excel-də başlıq sətri tapılmadı.")
        labels = [_header(v) if v is not None else "" for v in headings]
        found: dict[str, int] = {}
        for field, allowed in _FIELDS.items():
            positions = [i for i, value in enumerate(labels) if value in allowed]
            if len(positions) != 1:
                raise SaleValidationError(
                    f"Başlıq çatışmır və ya təkrarlanır: {field}. "
                    "Müştəri, Tarix, Faktura, Məbləğ, Valyuta sütunları lazımdır."
                )
            found[field] = positions[0]
        draft_rows: list[SaleDraft] = []
        invoices: set[tuple[str, str]] = set()
        for row_number, cells in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in cells):
                continue
            if len(draft_rows) >= MAX_INVOICES:
                raise SaleValidationError(f"Maksimum {MAX_INVOICES} faktura idxal edilə bilər.")
            try:
                customer_raw = cells[found["customer"]]
                date_raw = cells[found["date"]]
                invoice_raw = cells[found["invoice"]]
                amount_raw = cells[found["amount"]]
                currency_raw = cells[found["currency"]]
            except IndexError as error:
                raise SaleValidationError(f"{row_number}-ci sətr natamamdır.") from error
            if (
                not isinstance(customer_raw, str)
                or not isinstance(invoice_raw, str)
                or not isinstance(currency_raw, str)
            ):
                raise SaleValidationError(
                    f"{row_number}-ci sətr: müştəri, faktura və valyuta mətn olmalıdır."
                )
            customer = customer_raw.strip()
            invoice = invoice_raw.strip()
            currency = currency_raw.strip().upper()
            if not (0 < len(customer) <= 255 and 0 < len(invoice) <= 100):
                raise SaleValidationError(f"{row_number}-ci sətr: müştəri/faktura düzgün deyil.")
            if len(currency) != 3 or not currency.isascii() or not currency.isalpha():
                raise SaleValidationError(f"{row_number}-ci sətr: valyuta üçhərfli kod olmalıdır.")
            sold_on = _date(date_raw, row_number)
            if sold_on > (today or date.today()):
                raise SaleValidationError(f"{row_number}-ci sətr: gələcək satış qəbul edilmir.")
            if not isinstance(amount_raw, (str, int, float, Decimal)):
                raise SaleValidationError(f"{row_number}-ci sətr: məbləğ düzgün deyil.")
            if isinstance(amount_raw, bool) or (
                isinstance(amount_raw, str) and amount_raw.strip().startswith("=")
            ):
                raise SaleValidationError(f"{row_number}-ci sətr: formula qəbul edilmir.")
            try:
                amount = Decimal(str(amount_raw).strip())
            except InvalidOperation as error:
                raise SaleValidationError(f"{row_number}-ci sətr: məbləğ düzgün deyil.") from error
            exponent = amount.as_tuple().exponent
            if (
                not amount.is_finite()
                or amount <= 0
                or amount > Decimal("9999999999999999.99")
                or not isinstance(exponent, int)
                or exponent < -2
            ):
                raise SaleValidationError(f"{row_number}-ci sətr: məbləğ qeyri-məqbuldur.")
            key = (" ".join(customer.casefold().split()), invoice)
            if key in invoices:
                raise SaleValidationError(
                    f"{row_number}-ci sətr: eyni müştəri və faktura təkrarlandı."
                )
            invoices.add(key)
            draft_rows.append(SaleDraft(customer, sold_on, invoice, amount, currency))
        if not draft_rows:
            raise SaleValidationError("Heç bir etibarlı faktura sətri tapılmadı.")
        return tuple(draft_rows)
    finally:
        workbook.close()
