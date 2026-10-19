"""Bounded, formula-free stock XLSX preview; no data writes before approval."""

from dataclasses import dataclass, replace
from hashlib import sha256
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from office_assistant.application.inventory import (
    InventoryLedgerService,
    InventoryValidationError,
    StockDraft,
    StockEntry,
    parse_stock_draft,
)
from office_assistant.application.inventory_registry import VALID_UNITS
from office_assistant.infrastructure.database.models import User

MAX_XLSX_BYTES = 5 * 1024 * 1024
MAX_XLSX_UNPACKED = 35 * 1024 * 1024
MAX_ROWS = 200
_HEADERS = ("product_code", "warehouse", "kind", "quantity", "reference")
_UNIT_HEADERS = ("product_code", "warehouse", "kind", "quantity", "unit", "reference")


@dataclass(frozen=True, slots=True)
class StockXlsxPreview:
    sha256: str
    drafts: tuple[StockDraft, ...]
    opening_count: int
    receipt_count: int
    issue_count: int

    @property
    def count(self) -> int:
        return len(self.drafts)

    def safe_summary(self) -> dict[str, str | int]:
        """Never print private SKU/location/reference strings as log telemetry."""
        return {
            "sha256": self.sha256,
            "rows": self.count,
            "openings": self.opening_count,
            "receipts": self.receipt_count,
            "issues": self.issue_count,
        }


def _check_zip(payload: bytes) -> None:
    if not 0 < len(payload) <= MAX_XLSX_BYTES:
        raise InventoryValidationError("Excel faylı ölçü limitini aşır.")
    try:
        with ZipFile(BytesIO(payload)) as archive:
            contents = archive.infolist()
            if len(contents) > 400 or sum(x.file_size for x in contents) > MAX_XLSX_UNPACKED:
                raise InventoryValidationError("Excel arxivinin həcmi təhlükəsiz deyil.")
            for part in contents:
                name = part.filename.casefold().replace("\\", "/")
                if any(x in name for x in ("externallinks/", "vbaproject.bin", "embeddings/")):
                    raise InventoryValidationError("Excel xarici əlaqə və ya makro saxlaya bilməz.")
    except (BadZipFile, OSError) as error:
        raise InventoryValidationError("Excel faylı zədələnib.") from error


def preview_stock_xlsx(payload: bytes) -> StockXlsxPreview:
    """Parse the supplied bytes only; the caller must recheck SHA before commit."""
    _check_zip(payload)
    try:
        workbook = load_workbook(
            BytesIO(payload), read_only=True, data_only=False, keep_links=False
        )
    except (OSError, ValueError, InvalidFileException, BadZipFile) as error:
        raise InventoryValidationError("Excel sənədi açıla bilmədi.") from error
    try:
        if len(workbook.worksheets) != 1:
            raise InventoryValidationError("Yalnız bir vərəqli stok şablonu qəbul edilir.")
        sheet = workbook.worksheets[0]
        if sheet.max_column not in (len(_HEADERS), len(_UNIT_HEADERS)) or sheet.max_row > MAX_ROWS + 1:
            raise InventoryValidationError("Excel sütunları və ya sətir limiti düzgün deyil.")
        rows = sheet.iter_rows()
        header = next(rows, ())
        names = tuple(str(cell.value or "").strip().casefold() for cell in header)
        if names not in (_HEADERS, _UNIT_HEADERS):
            raise InventoryValidationError("Excel başlıqları təsdiqlənmiş şablona uyğun deyil.")
        drafts: list[StockDraft] = []
        seen: set[tuple[str, str, str]] = set()
        for row in rows:
            if len(drafts) >= MAX_ROWS:
                raise InventoryValidationError("Excel sətir limiti aşılıb.")
            if len(row) != len(names):
                raise InventoryValidationError("Excel sətrində sütun çatışmır.")
            if any(
                cell.data_type == "f"
                or (isinstance(cell.value, str) and cell.value.lstrip().startswith("="))
                for cell in row
            ):
                raise InventoryValidationError("Excel düsturuna icazə verilmir.")
            values = [str(cell.value if cell.value is not None else "").strip() for cell in row]
            if all(not value for value in values):
                raise InventoryValidationError("Excel daxilində boş sətir olmamalıdır.")
            if any(
                not value or "|" in value or any(ord(ch) < 32 for ch in value) for value in values
            ):
                raise InventoryValidationError("Excel sətrində yararsız məlumat var.")
            unit = None
            if names == _UNIT_HEADERS:
                unit = values[4].upper()
                if unit not in VALID_UNITS:
                    raise InventoryValidationError("Excel ölçü vahidi təsdiqlənməyib.")
                values = [*values[:4], values[5]]
            text = "Anbar əməliyyatı: " + " | ".join(values)
            draft = parse_stock_draft(text)
            if draft is None:
                raise InventoryValidationError("Excel sətri tanınmadı.")
            if unit is not None:
                draft = replace(draft, unit=unit)
            key = (
                draft.product_code,
                " ".join(draft.warehouse.casefold().split()),
                " ".join(draft.reference.casefold().split()),
            )
            if key in seen:
                raise InventoryValidationError("Eyni stok sənədi Excel daxilində təkrarlanır.")
            seen.add(key)
            drafts.append(draft)
        if not drafts:
            raise InventoryValidationError("Excel boşdur, yazılacaq əməliyyat yoxdur.")
        return StockXlsxPreview(
            sha256(payload).hexdigest(),
            tuple(drafts),
            sum(row.kind == "OPENING" for row in drafts),
            sum(row.kind == "RECEIPT" for row in drafts),
            sum(row.kind == "ISSUE" for row in drafts),
        )
    finally:
        workbook.close()


def approved_drafts(payload: bytes, *, confirmed_sha256: str) -> tuple[StockDraft, ...]:
    """Return content only for exact, intentionally confirmed workbook bytes."""
    import hmac

    preview = preview_stock_xlsx(payload)
    if not hmac.compare_digest(preview.sha256, confirmed_sha256):
        raise InventoryValidationError("Önbaxışdan sonra Excel dəyişib; təsdiq etibarsızdır.")
    return preview.drafts


async def commit_stock_xlsx(
    ledger: InventoryLedgerService,
    actor: User,
    payload: bytes,
    *,
    confirmed_sha256: str,
) -> tuple[StockEntry, ...]:
    """Only an explicitly confirmed, SHA-identical workbook may post stock."""
    drafts = approved_drafts(payload, confirmed_sha256=confirmed_sha256)
    return await ledger.record_batch(actor, drafts, workbook_sha256=confirmed_sha256)
