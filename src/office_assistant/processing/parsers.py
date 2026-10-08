"""Targeted, deterministic parser adapters and their registry."""

# Third-party document libraries do not all ship complete Pyright stubs. The
# adapter boundary keeps that uncertainty out of the normalized model/worker.
# pyright: reportMissingTypeStubs=false, reportUnknownMemberType=false, reportUnknownVariableType=false, reportArgumentType=false, reportOptionalMemberAccess=false

import zipfile
from pathlib import Path
from posixpath import normpath
from typing import Any, Protocol, cast
from xml.etree import ElementTree

import pypdf
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.cell.cell import TYPE_FORMULA
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pypdf import PdfReader
from pypdf.generic import ContentStream

from office_assistant.infrastructure.database.models import ImageAssetOrigin
from office_assistant.processing.image_metadata import image_signature
from office_assistant.processing.models import ContentBlock, EmbeddedImage, ParsedDocument
from office_assistant.processing.pdf_images import (
    SampleImageError,
    pdf_name,
    picture_bytes,
    soft_mask_note,
)

# A picture inside a document is untrusted input like every other byte the bot reads, so extraction
# is bounded: one payload and one count cap keep a hostile file from turning a parse into a disk
# write marathon. Both caps are ASSUMPTION values (Master Spec §23 leaves them open, Q29).
MAX_EMBEDDED_IMAGE_BYTES = 8 * 1024 * 1024
MAX_EMBEDDED_IMAGES_PER_DOCUMENT = 200

_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"
_SPREADSHEET_DRAWING_NAMESPACE = (
    "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
)
_DRAWML_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/main"
_SPREADSHEETML_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_XML_NAMESPACES = {"a": _DRAWML_NAMESPACE, "xdr": _SPREADSHEET_DRAWING_NAMESPACE}
_BLIP_TAG = qn("a:blip")
_RELATIONSHIP_ID_ATTRIBUTE = qn("r:embed")
_RELATIONSHIP_ID_FIELD = f"{{{_RELATIONSHIP_NAMESPACE}}}id"
_SHEET_TAG = f"{{{_SPREADSHEETML_NAMESPACE}}}sheet"
_RELATIONSHIP_TAG = f"{{{_PACKAGE_RELATIONSHIP_NAMESPACE}}}Relationship"
_WORKSHEET_RELATIONSHIP = f"{_RELATIONSHIP_NAMESPACE}/worksheet"
_DRAWING_RELATIONSHIP = f"{_RELATIONSHIP_NAMESPACE}/drawing"
_IMAGE_RELATIONSHIP = f"{_RELATIONSHIP_NAMESPACE}/image"

# The six header and footer roles one Word section can define, as the attribute that reaches them
# and the role a locator should name them by.
_SECTION_DEFINITIONS = (
    ("header", "default_header"),
    ("first_page_header", "first_page_header"),
    ("even_page_header", "even_page_header"),
    ("footer", "default_footer"),
    ("first_page_footer", "first_page_footer"),
    ("even_page_footer", "even_page_footer"),
)


class _ExtractedImages:
    """One document's pictures, bounded, plus an honest record of every payload left behind."""

    def __init__(self, warnings: list[str]) -> None:
        self.images: list[EmbeddedImage] = []
        self.warnings = warnings
        self._stopped = False

    def add(self, locator: dict[str, Any], data: bytes, description: str) -> bool:
        """Keep one picture, or say out loud why it was left behind. Returns whether it was kept."""

        if len(self.images) >= MAX_EMBEDDED_IMAGES_PER_DOCUMENT:
            if not self._stopped:
                self._stopped = True
                self.warnings.append(
                    f"Picture extraction stopped at {MAX_EMBEDDED_IMAGES_PER_DOCUMENT} images"
                )
            return False
        if len(data) > MAX_EMBEDDED_IMAGE_BYTES:
            self.warnings.append(
                f"{description}: {len(data)} bytes exceeds the {MAX_EMBEDDED_IMAGE_BYTES} byte cap"
            )
            return False
        signature = image_signature(data)
        if signature is None:
            self.warnings.append(
                f"{description}: bytes are not a PNG or JPEG this build can describe"
            )
            return False
        self.images.append(EmbeddedImage(locator, data, signature.mime_type))
        return True


class DocumentParser(Protocol):
    name: str
    version: str
    #: Provenance recorded for pictures found inside this format, or ``None`` when the format
    #: carries none. The standalone upload of an image is registered by the worker, not here.
    embedded_origin: ImageAssetOrigin | None

    def parse(self, path: Path, *, source_file_id: int, media_type: str) -> ParsedDocument: ...


class TxtParser:
    name = "txt"
    version = "1.0"
    embedded_origin: ImageAssetOrigin | None = None

    def parse(self, path: Path, *, source_file_id: int, media_type: str) -> ParsedDocument:
        content = path.read_text(encoding="utf-8-sig")
        blocks = (
            [
                ContentBlock(
                    sequence=1,
                    block_type="plain_text",
                    text=content,
                    locator={"line_start": 1, "line_end": max(1, len(content.splitlines()))},
                )
            ]
            if content
            else []
        )
        return ParsedDocument(source_file_id, self.name, self.version, media_type, {}, blocks)


class PdfParser:
    name = "pypdf"
    version = pypdf.__version__
    embedded_origin: ImageAssetOrigin | None = ImageAssetOrigin.PDF_EMBEDDED

    def parse(self, path: Path, *, source_file_id: int, media_type: str) -> ParsedDocument:
        blocks: list[ContentBlock] = []
        text_warnings: list[str] = []
        reader = PdfReader(str(path), strict=True)
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                blocks.append(
                    ContentBlock(
                        len(blocks) + 1,
                        "page_text",
                        str(cast(Any, text)),
                        {"page": page_number},
                    )
                )
            else:
                text_warnings.append(
                    f"Page {page_number} contains no extractable text; OCR may be required"
                )
        image_warnings: list[str] = []
        images = self._embedded_images(reader, image_warnings)
        return ParsedDocument(
            source_file_id,
            self.name,
            self.version,
            media_type,
            {},
            blocks,
            [*text_warnings, *image_warnings],
            requires_ocr=bool(text_warnings),
            embedded_images=images,
        )

    def _embedded_images(self, reader: PdfReader, warnings: list[str]) -> list[EmbeddedImage]:
        """Pictures the page content stream actually draws, in draw order.

        Only drawn XObjects are recorded: a page resource can name a picture the page never paints,
        and an occurrence claiming such a page would be an invented locator (§23.5). The page's
        ``image_index`` counts every drawn picture, including one this build cannot store, so the
        number a run records does not move when a reader learns a new format.
        """

        collected = _ExtractedImages(warnings)
        for page_number, page in enumerate(reader.pages, start=1):
            operations = self._page_operations(page, page_number, warnings)
            if operations is None:
                continue
            inline = sum(1 for _, operator in operations if operator == b"INLINE IMAGE")
            if inline:
                warnings.append(
                    f"Page {page_number}: {inline} inline image(s) are not stored as picture assets"
                )
            resources = _xobject_resources(page)
            index = 0
            for name in _drawn_names(operations):
                stream = resources.get(name)
                if stream is None:
                    continue
                subtype = pdf_name(stream.get("/Subtype"))
                if subtype == "/Image":
                    index += 1
                    self._collect(
                        collected,
                        stream,
                        _page_locator(page_number, index, name),
                        f"Page {page_number} picture {name}",
                    )
                elif subtype == "/Form":
                    index = self._collect_form_images(
                        collected, reader, stream, page_number, name, index
                    )
        return collected.images

    def _collect_form_images(
        self,
        collected: _ExtractedImages,
        reader: PdfReader,
        form: Any,
        page_number: int,
        form_name: str,
        index: int,
    ) -> int:
        """One Form level deep: the form's own draw order says which pictures it paints."""

        try:
            operations = list(ContentStream(form, reader).operations)
        except Exception as error:  # noqa: BLE001 - a damaged form must not cost the whole page
            collected.warnings.append(
                f"Page {page_number} form {form_name}: content unreadable ({type(error).__name__}),"
                " its pictures are not recorded"
            )
            return index
        resources = _xobject_resources(form)
        for name in _drawn_names(operations):
            stream = resources.get(name)
            if stream is None or pdf_name(stream.get("/Subtype")) != "/Image":
                continue
            index += 1
            locator = _page_locator(page_number, index, form_name)
            locator["form_xobject"] = name
            description = f"Page {page_number} picture {form_name}{name}"
            self._collect(collected, stream, locator, description)
        return index

    @staticmethod
    def _collect(
        collected: _ExtractedImages,
        stream: Any,
        locator: dict[str, Any],
        description: str,
    ) -> None:
        try:
            data = picture_bytes(stream, MAX_EMBEDDED_IMAGE_BYTES)
        except SampleImageError as error:
            collected.warnings.append(f"{description}: {error}")
            return
        # A mask is worth naming only once a picture was kept: if nothing was stored, no stored
        # picture left a mask behind.
        note = soft_mask_note(stream) if collected.add(locator, data, description) else None
        if note is not None:
            collected.warnings.append(f"{description}: {note}")

    @staticmethod
    def _page_operations(
        page: Any, page_number: int, warnings: list[str]
    ) -> list[tuple[Any, bytes]] | None:
        try:
            contents = page.get_contents()
            return list(contents.operations) if contents is not None else []
        except Exception as error:  # noqa: BLE001 - an unreadable page must not fail the document
            warnings.append(
                f"Page {page_number}: content stream unreadable ({type(error).__name__}),"
                " its pictures are not recorded"
            )
            return None


def _drawn_names(operations: list[tuple[Any, bytes]]) -> list[str]:
    names: list[str] = []
    for operands, operator in operations:
        if operator == b"Do" and operands:
            names.append(pdf_name(operands[0]))
    return [name for name in names if name]


def _xobject_resources(holder: Any) -> dict[str, Any]:
    resources = holder.get_inherited("/Resources") if hasattr(holder, "get_inherited") else None
    if resources is None:
        resources = holder.get("/Resources")
    if resources is None:
        return {}
    resolved = resources.get_object() if hasattr(resources, "get_object") else resources
    xobjects = resolved.get("/XObject")
    if xobjects is None:
        return {}
    xobjects = xobjects.get_object() if hasattr(xobjects, "get_object") else xobjects
    if not hasattr(xobjects, "items"):
        return {}
    return {
        str(key): (value.get_object() if hasattr(value, "get_object") else value)
        for key, value in xobjects.items()
    }


def _page_locator(page_number: int, image_index: int, name: str) -> dict[str, Any]:
    return {"page": page_number, "image_index": image_index, "xobject": name}


def _zip_root(archive: zipfile.ZipFile, name: str) -> ElementTree.Element | None:
    """Parse one XML member of a package, or report nothing for a member that is absent."""

    if name not in archive.namelist():
        return None
    try:
        return ElementTree.fromstring(archive.read(name))
    except ElementTree.ParseError:
        return None


def _relationships(archive: zipfile.ZipFile, source_part: str, kind: str) -> list[tuple[str, str]]:
    """``(id, target)`` for every relationship of one type held by a package part."""

    directory, _, leaf = source_part.rpartition("/")
    root = _zip_root(archive, f"{directory}/_rels/{leaf}.rels")
    if root is None:
        return []
    return [
        (
            relationship.get("Id") or "",
            relationship.get("Target") or "",
        )
        for relationship in root.iter(_RELATIONSHIP_TAG)
        if relationship.get("Type") == kind
    ]


def _resolve_part(source_part: str, target: str) -> str:
    """Turn a relationship target into a package member name, absolute or part-relative."""

    if target.startswith("/"):
        return normpath(target.removeprefix("/"))
    directory, _, _ = source_part.rpartition("/")
    return normpath(f"{directory}/{target}") if directory else normpath(target)


def _zip_payload(
    archive: zipfile.ZipFile, name: str, description: str, warnings: list[str]
) -> bytes | None:
    """Read one member, refusing to decompress more than a picture could ever be."""

    try:
        declared_size = archive.getinfo(name).file_size
    except KeyError:
        warnings.append(f"{description}: member {name} is absent from this package")
        return None
    if declared_size > MAX_EMBEDDED_IMAGE_BYTES:
        warnings.append(
            f"{description}: {declared_size} declared bytes exceeds the {MAX_EMBEDDED_IMAGE_BYTES}"
            " byte cap"
        )
        return None
    return archive.read(name)


def _anchor_cell(anchor: ElementTree.Element) -> str | None:
    """The cell a drawing anchor sits on, or ``None`` when the anchor declares no cell."""

    origin = anchor.find("xdr:from", _XML_NAMESPACES)
    if origin is None:
        return None
    column = origin.find("xdr:col", _XML_NAMESPACES)
    row = origin.find("xdr:row", _XML_NAMESPACES)
    if column is None or row is None or not (column.text or "").isdigit():
        return None
    if not (row.text or "").isdigit():
        return None
    return f"{get_column_letter(int(column.text) + 1)}{int(row.text) + 1}"


class DocxParser:
    name = "python-docx"
    version = "1.2.0"
    embedded_origin: ImageAssetOrigin | None = ImageAssetOrigin.DOCX_EMBEDDED

    def parse(self, path: Path, *, source_file_id: int, media_type: str) -> ParsedDocument:
        document = Document(str(path))
        blocks: list[ContentBlock] = []
        warnings: list[str] = []
        collected = _ExtractedImages(warnings)
        self._collect_block_images(collected, document.part, document.element.body, {}, "body")
        self._collect_section_images(collected, document)
        paragraph_index = 0
        table_index = 0
        for element in document.element.body.iterchildren():
            if element.tag.endswith("}p"):
                paragraph = Paragraph(cast(Any, element), document)
                current_paragraph_index = paragraph_index
                paragraph_index += 1
                text = paragraph.text.strip()
                if text:
                    style = paragraph.style.name if paragraph.style else ""
                    block_type = "heading" if style.lower().startswith("heading") else "paragraph"
                    blocks.append(
                        ContentBlock(
                            len(blocks) + 1,
                            block_type,
                            text,
                            {"paragraph_index": current_paragraph_index},
                        )
                    )
            elif element.tag.endswith("}tbl"):
                table = Table(cast(Any, element), document)
                current_table_index = table_index
                table_index += 1
                for row_index, row in enumerate(table.rows, start=1):
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        blocks.append(
                            ContentBlock(
                                len(blocks) + 1,
                                "table_row",
                                " | ".join(values),
                                {
                                    "table_index": current_table_index,
                                    "row": row_index,
                                },
                                {"cells": values},
                            )
                        )
        return ParsedDocument(
            source_file_id,
            self.name,
            self.version,
            media_type,
            {},
            blocks,
            warnings,
            embedded_images=collected.images,
        )

    @classmethod
    def _collect_section_images(cls, collected: _ExtractedImages, document: Any) -> None:
        """Pictures a section's own header or footer carries, located in the part that holds them.

        A header is a package part, not a place in the body, so its locator names the part and the
        role it plays rather than pretending to be body text. A section that links its header to the
        previous one owns no definition and therefore no picture of its own: the section that does
        define it has already been read, and a second sighting invented here would be a duplicate
        claim (§23.5).
        """

        for section_index, section in enumerate(document.sections):
            for attribute, role in _SECTION_DEFINITIONS:
                holder = getattr(section, attribute)
                if holder.is_linked_to_previous:
                    continue
                part = holder.part
                cls._collect_block_images(
                    collected,
                    part,
                    part.element,
                    {
                        "part": str(part.partname).lstrip("/"),
                        "section_index": section_index,
                        "role": role,
                    },
                    f"{role} of section {section_index + 1}",
                )

    @staticmethod
    def _collect_block_images(
        collected: _ExtractedImages,
        part: Any,
        root: Any,
        context: dict[str, Any],
        description_prefix: str,
    ) -> None:
        """Every picture in a run of block elements, each at the locator of its own place.

        Rows and cells come from the XML rather than ``table.rows[i].cells[j]``: a merged cell is
        reported once per grid position it spans, which would let one picture claim several homes.
        """

        paragraph_index = 0
        table_index = 0
        for element in root.iterchildren():
            if element.tag.endswith("}p"):
                current_paragraph_index = paragraph_index
                paragraph_index += 1
                DocxParser._collect_images(
                    collected,
                    part,
                    element,
                    {**context, "paragraph_index": current_paragraph_index},
                    f"{description_prefix} paragraph {current_paragraph_index}",
                )
            elif element.tag.endswith("}tbl"):
                current_table_index = table_index
                table_index += 1
                for row_index, row_element in enumerate(element.findall(qn("w:tr")), start=1):
                    for column_index, cell_element in enumerate(row_element.findall(qn("w:tc")), 1):
                        DocxParser._collect_images(
                            collected,
                            part,
                            cell_element,
                            {
                                **context,
                                "table_index": current_table_index,
                                "row": row_index,
                                "column": column_index,
                            },
                            f"{description_prefix} table {current_table_index} "
                            f"cell {row_index}:{column_index}",
                        )

    @staticmethod
    def _collect_images(
        collected: _ExtractedImages,
        part: Any,
        element: Any,
        context: dict[str, Any],
        description: str,
    ) -> None:
        """Every picture this element draws, in document order, through its relationship id."""

        related_parts = part.related_parts
        for image_index, blip in enumerate(element.iter(_BLIP_TAG), start=1):
            relationship_id = blip.get(_RELATIONSHIP_ID_ATTRIBUTE)
            if not relationship_id or relationship_id not in related_parts:
                collected.warnings.append(
                    f"{description} picture {image_index}: relationship "
                    f"{relationship_id or 'none'} is absent from this package"
                )
                continue
            locator = {**context, "image_index": image_index, "relationship_id": relationship_id}
            collected.add(
                locator,
                related_parts[relationship_id].blob,
                f"{description} picture {image_index}",
            )


class XlsxParser:
    name = "openpyxl"
    version = "3.1.5"
    embedded_origin: ImageAssetOrigin | None = ImageAssetOrigin.XLSX_EMBEDDED

    def parse(self, path: Path, *, source_file_id: int, media_type: str) -> ParsedDocument:
        workbook = load_workbook(path, data_only=False)
        blocks: list[ContentBlock] = []
        try:
            for sheet in workbook.worksheets:
                self._parse_sheet(sheet, blocks)
        finally:
            workbook.close()
        warnings: list[str] = []
        collected = _ExtractedImages(warnings)
        self._collect_images(path, collected)
        return ParsedDocument(
            source_file_id,
            self.name,
            self.version,
            media_type,
            {},
            blocks,
            warnings,
            embedded_images=collected.images,
        )

    @classmethod
    def _collect_images(cls, path: Path, collected: _ExtractedImages) -> None:
        """Pictures anchored on a worksheet, read from the package rather than a live workbook.

        openpyxl's drawing reader returns no pictures at all when Pillow is absent
        (``reader.drawings.find_images`` gives up before it builds one), and Pillow is not a
        dependency of this build (Q28). Following the drawing parts directly also reads the anchor
        the drawing itself declares, which is what §23.5 asks a locator to name.
        """

        with zipfile.ZipFile(path) as archive:
            for sheet_name, sheet_part in cls._sheets(archive, collected):
                drawings: list[str] = []
                for _, target in _relationships(archive, sheet_part, _DRAWING_RELATIONSHIP):
                    part = _resolve_part(sheet_part, target)
                    if part not in drawings:
                        drawings.append(part)
                for drawing_part in drawings:
                    root = _zip_root(archive, drawing_part)
                    if root is None:
                        collected.warnings.append(
                            f"Sheet {sheet_name}: drawing part {drawing_part} is unreadable"
                        )
                        continue
                    relationships = {
                        relationship_id: _resolve_part(drawing_part, target)
                        for relationship_id, target in _relationships(
                            archive, drawing_part, _IMAGE_RELATIONSHIP
                        )
                    }
                    image_index = 0
                    for anchor in root:
                        anchor_cell = _anchor_cell(anchor)
                        for blip in anchor.iter(_BLIP_TAG):
                            image_index += 1
                            relationship_id = blip.get(_RELATIONSHIP_ID_ATTRIBUTE)
                            member = relationships.get(relationship_id) if relationship_id else None
                            description = f"Sheet {sheet_name} picture {image_index}"
                            if member is None:
                                collected.warnings.append(
                                    f"{description}: relationship "
                                    f"{relationship_id or 'none'} is absent from this package"
                                )
                                continue
                            locator: dict[str, Any] = {
                                "sheet": sheet_name,
                                "image_index": image_index,
                                "relationship_id": relationship_id,
                            }
                            if anchor_cell is not None:
                                locator["anchor_cell"] = anchor_cell
                            data = _zip_payload(archive, member, description, collected.warnings)
                            if data is not None:
                                collected.add(locator, data, description)

    @staticmethod
    def _sheets(archive: zipfile.ZipFile, collected: _ExtractedImages) -> list[tuple[str, str]]:
        """Worksheet names with their part paths, in the order the workbook lists them."""

        root = _zip_root(archive, "xl/workbook.xml")
        if root is None:
            collected.warnings.append("Workbook part xl/workbook.xml is unreadable")
            return []
        names = {
            sheet.get(_RELATIONSHIP_ID_FIELD): (sheet.get("name") or "")
            for sheet in root.iter(_SHEET_TAG)
        }
        sheets: list[tuple[str, str]] = []
        workbook_rels = _relationships(archive, "xl/workbook.xml", _WORKSHEET_RELATIONSHIP)
        for relationship_id, target in workbook_rels:
            name = names.get(relationship_id)
            if name is None:
                collected.warnings.append(
                    f"Worksheet relationship {relationship_id} has no sheet entry"
                )
                continue
            sheets.append((name, _resolve_part("xl/workbook.xml", target)))
        return sheets

    @staticmethod
    def _parse_sheet(sheet: Worksheet, blocks: list[ContentBlock]) -> None:
        for row in sheet.iter_rows():
            non_empty = [cell for cell in row if cell.value is not None]
            if not non_empty:
                continue
            values = [cell.value for cell in non_empty]
            formulas = [cell.data_type == TYPE_FORMULA for cell in non_empty]
            start = non_empty[0].coordinate
            end = non_empty[-1].coordinate
            blocks.append(
                ContentBlock(
                    len(blocks) + 1,
                    "table_row",
                    " | ".join(str(value) for value in values),
                    {"sheet": sheet.title, "range": f"{start}:{end}"},
                    {"values": values, "formulas": formulas},
                )
            )


class ImageProbeParser:
    """Verify uploaded image bytes before a vision call is allowed to spend them.

    The header, not the extension, decides the media type: a file whose bytes are not a
    supported raster image fails the run instead of being sent to the vendor.
    """

    name = "image-probe"
    version = "1.0"
    embedded_origin: ImageAssetOrigin | None = None

    def parse(self, path: Path, *, source_file_id: int, media_type: str) -> ParsedDocument:
        detected = detect_image_media_type(path.read_bytes()[:12])
        if detected is None:
            raise ValueError("unsupported_image_signature")
        declared = media_type.split(";", 1)[0].lower()
        if declared not in {"", "application/octet-stream"} and declared != detected:
            raise ValueError("image_signature_mismatch")
        return ParsedDocument(
            source_file_id,
            self.name,
            self.version,
            detected,
            {"filename": path.name},
            [],
            [f"{detected} image requires vision extraction"],
            requires_vision=True,
        )


def detect_image_media_type(header: bytes) -> str | None:
    if header.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if header.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    return None


class ParserRegistry:
    def __init__(self) -> None:
        self._parsers: dict[str, DocumentParser] = {
            ".txt": TxtParser(),
            ".pdf": PdfParser(),
            ".docx": DocxParser(),
            ".xlsx": XlsxParser(),
            ".jpg": ImageProbeParser(),
            ".jpeg": ImageProbeParser(),
            ".png": ImageProbeParser(),
        }

    def for_filename(self, filename: str) -> DocumentParser:
        """Route by extension.

        `application/file_validation.validate_upload` already ties the extension to the declared
        content type, and the image parser re-proves the type from the bytes themselves, so a
        second content-type lookup here would only duplicate that check.
        """

        suffix = Path(filename).suffix.lower()
        try:
            return self._parsers[suffix]
        except KeyError as error:
            raise ValueError(f"No parser registered for {suffix}") from error
