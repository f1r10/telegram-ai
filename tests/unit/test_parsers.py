from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf.errors import PdfReadError

from office_assistant.processing.parsers import ParserRegistry

PNG_HEADER = b"\x89PNG\r\n\x1a\n"
JPEG_BODY = b"\xff\xd8\xff\xe0" + b"\x00" * 32


def test_txt_parser_preserves_line_provenance(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("first line\nsecond line\n", encoding="utf-8")

    parsed = (
        ParserRegistry()
        .for_filename("notes.txt")
        .parse(path, source_file_id=1, media_type="text/plain")
    )

    assert parsed.blocks[0].locator == {"line_start": 1, "line_end": 2}
    assert parsed.blocks[0].text == "first line\nsecond line\n"


def test_parsers_handle_empty_txt_and_images(tmp_path: Path) -> None:
    empty = tmp_path / "empty.txt"
    empty.write_text("", encoding="utf-8")
    image = tmp_path / "photo.png"
    image.write_bytes(PNG_HEADER + b"\x00" * 24)

    assert (
        ParserRegistry()
        .for_filename("empty.txt")
        .parse(empty, source_file_id=1, media_type="text/plain")
        .blocks
        == []
    )
    probed = (
        ParserRegistry()
        .for_filename("photo.png")
        .parse(image, source_file_id=2, media_type="image/png")
    )
    assert probed.requires_vision
    assert probed.blocks == []
    assert probed.parser_name == "image-probe"
    assert probed.warnings == ["image/png image requires vision extraction"]


def test_image_probe_trusts_bytes_over_the_declared_type(tmp_path: Path) -> None:
    jpeg = tmp_path / "photo.jpg"
    jpeg.write_bytes(JPEG_BODY)

    probed = (
        ParserRegistry()
        .for_filename("photo.jpg")
        .parse(jpeg, source_file_id=3, media_type="application/octet-stream")
    )

    assert probed.media_type == "image/jpeg"
    assert probed.requires_vision


@pytest.mark.parametrize(
    ("filename", "body", "error"),
    [
        ("photo.jpg", b"PROCBINDnotanimage", "unsupported_image_signature"),
        ("photo.png", PNG_HEADER[:6] + b"\x00" * 30, "unsupported_image_signature"),
        ("photo.jpg", JPEG_BODY, "image_signature_mismatch"),
    ],
)
def test_image_probe_refuses_bytes_that_are_not_a_supported_image(
    tmp_path: Path, filename: str, body: bytes, error: str
) -> None:
    path = tmp_path / filename
    path.write_bytes(body)
    declared = "image/png" if error == "image_signature_mismatch" else "image/jpeg"

    with pytest.raises(ValueError, match=error):
        ParserRegistry().for_filename(filename).parse(path, source_file_id=4, media_type=declared)


def test_pdf_parser_preserves_page_numbers(tmp_path: Path) -> None:
    path = tmp_path / "two-pages.pdf"
    path.write_bytes(_make_pdf(["page one", "page two"]))

    parsed = (
        ParserRegistry()
        .for_filename("two-pages.pdf")
        .parse(path, source_file_id=3, media_type="application/pdf")
    )

    assert [block.locator["page"] for block in parsed.blocks] == [1, 2]
    assert not parsed.requires_ocr


def test_pdf_parser_flags_empty_page_for_ocr(tmp_path: Path) -> None:
    path = tmp_path / "scanned-like.pdf"
    path.write_bytes(_make_pdf(["page one", None]))

    parsed = (
        ParserRegistry()
        .for_filename("scanned-like.pdf")
        .parse(path, source_file_id=6, media_type="application/pdf")
    )

    assert parsed.requires_ocr
    assert parsed.warnings == ["Page 2 contains no extractable text; OCR may be required"]


def test_pdf_parser_rejects_corrupt_pdf(tmp_path: Path) -> None:
    path = tmp_path / "corrupt.pdf"
    path.write_bytes(b"not a PDF")

    with pytest.raises(PdfReadError):
        ParserRegistry().for_filename("corrupt.pdf").parse(
            path, source_file_id=7, media_type="application/pdf"
        )


def test_docx_parser_preserves_blocks_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "document.docx"
    document = Document()
    document.add_heading("Heading", level=1)
    document.add_paragraph("Paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.save(str(path))

    blocks = (
        ParserRegistry()
        .for_filename("document.docx")
        .parse(
            path,
            source_file_id=4,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        .blocks
    )

    assert [block.block_type for block in blocks] == ["heading", "paragraph", "table_row"]
    assert blocks[-1].metadata["cells"] == ["A", "B"]


def test_xlsx_parser_skips_blank_regions_and_keeps_formulas(tmp_path: Path) -> None:
    path = tmp_path / "sheet.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Products"
    sheet["A1"] = "KA45"
    sheet["B1"] = 125
    sheet["C1"] = "=B1*2"
    workbook.create_sheet("Empty")
    workbook.save(str(path))
    workbook.close()

    parsed = (
        ParserRegistry()
        .for_filename("sheet.xlsx")
        .parse(
            path,
            source_file_id=5,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    )

    assert len(parsed.blocks) == 1
    assert parsed.blocks[0].locator == {"sheet": "Products", "range": "A1:C1"}
    assert parsed.blocks[0].metadata["formulas"] == [False, False, True]


def _make_pdf(pages: list[str | None]) -> bytes:
    objects: list[bytes] = []
    page_ids: list[int] = []
    for page_index, text in enumerate(pages):
        page_id = 3 + page_index * 2
        content_id = page_id + 1
        page_ids.append(page_id)
        stream = b"q Q" if text is None else f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.extend(
            [
                (
                    f"{page_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Contents {content_id} 0 R /Resources << /Font << "
                    f"/F1 {3 + len(pages) * 2} 0 R >> >> >> endobj"
                ).encode(),
                f"{content_id} 0 obj << /Length {len(stream)} >> stream\n".encode()
                + stream
                + b"\nendstream endobj",
            ]
        )
    font_id = 3 + len(pages) * 2
    prefix = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj",
        (
            f"2 0 obj << /Type /Pages /Kids "
            f"[{' '.join(f'{page_id} 0 R' for page_id in page_ids)}] "
            f"/Count {len(page_ids)} >> endobj"
        ).encode(),
    ]
    objects = (
        prefix
        + objects
        + [
            (
                f"{font_id} 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj"
            ).encode()
        ]
    )
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(output))
        output.extend(obj + b"\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        (
            f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(output)
