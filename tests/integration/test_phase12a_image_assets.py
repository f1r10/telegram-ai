"""Phase 12A: image content is stored once, every sighting keeps its own locator (§23.3-§23.5)."""

import asyncio
import shutil
from collections.abc import Awaitable, Callable
from datetime import UTC, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any, cast

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from PIL import Image, ImageDraw
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from office_assistant.core.settings import Settings, load_settings
from office_assistant.infrastructure.database.engine import create_database_engine
from office_assistant.infrastructure.database.models import (
    AuditEvent,
    DocumentParseRun,
    DocumentProcessingStatus,
    FileObject,
    FileObjectState,
    ImageAsset,
    ImageAssetOccurrence,
    ImageAssetOrigin,
    IngestionItemType,
    IngestionJob,
    IngestionJobStatus,
    IngestionSession,
    IngestionSessionItem,
    IngestionSessionStatus,
    KnowledgeIndexJob,
    OutboxEvent,
    PendingProductImageConfirmation,
    ProductImageRelation,
    SystemMetadata,
    User,
    UserRole,
    UserStatus,
)
from office_assistant.infrastructure.database.session import create_session_factory
from office_assistant.infrastructure.storage.local import LocalFileStorage
from office_assistant.internal_evidence.delivery import FileDeliveryError, SourceFileDeliveryService
from office_assistant.processing.document_visual_candidates import active_document_images
from office_assistant.processing.image_assets import ImageAssetRegistry
from office_assistant.processing.semantic_visual import CLIP_DIMENSION, SemanticVisualSearch
from office_assistant.processing.visual_ann import VisualANNIndex
from office_assistant.processing.worker_service import IngestionJobService

pytestmark = pytest.mark.integration

PNG_START = b"\x89PNG\r\n\x1a\n"
JPEG_BYTES = b"\xff\xd8\xff\xe0" + b"\x00" * 128
TELEGRAM_USER_ID = 941201


def _png(width: int = 240, height: int = 100) -> bytes:
    return (
        PNG_START
        + (13).to_bytes(4, "big")
        + b"IHDR"
        + width.to_bytes(4, "big")
        + height.to_bytes(4, "big")
        + b"\x00" * 8
    )


def _digest(data: bytes) -> str:
    return sha256(data).hexdigest()


def _derived_image(storage: LocalFileStorage, file_id: int, run_id: int, data: bytes) -> Path:
    """Write image bytes outside ``blobs/``, the way a derived-artifact writer would (§23.3)."""

    directory = storage.derived_root / str(file_id) / str(run_id)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / "image_0.png"
    path.write_bytes(data)
    return path


DRAWINGML = "http://schemas.openxmlformats.org/drawingml/2006/main"
RELATIONSHIPS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _docx_bytes(logo: bytes) -> bytes:
    """A real .docx that shows one logo twice: in a paragraph and in a table cell."""

    document = Document()
    document.add_paragraph("Qiymət siyahısı")
    document.add_table(rows=1, cols=2)
    part = Part(PackURI("/word/media/image1.png"), "image/png", logo, document.part.package)
    relationship_id = document.part.relate_to(part, RT.IMAGE)

    body = cast(Any, document.element).body
    body.findall(qn("w:p"))[0].append(_inline_picture(relationship_id))
    cell = body.findall(qn("w:tbl"))[0].findall(qn("w:tr"))[0].findall(qn("w:tc"))[1]
    cell.findall(qn("w:p"))[0].append(_inline_picture(relationship_id))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _inline_picture(relationship_id: str) -> Any:
    return parse_xml(
        "<w:drawing "
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        f'xmlns:a="{DRAWINGML}" xmlns:r="{RELATIONSHIPS}">'
        "<wp:inline><a:graphic><a:graphicData "
        f'uri="{DRAWINGML}/picture">'
        f'<a:blip r:embed="{relationship_id}"/>'
        "</a:graphicData></a:graphic></wp:inline></w:drawing>"
    )


def _settings_or_skip() -> Settings:
    settings = load_settings()
    if not all(
        (
            settings.postgres_db,
            settings.postgres_user,
            settings.postgres_password,
            settings.postgres_host,
        )
    ):
        pytest.skip("PostgreSQL integration settings are not configured")
    return settings


async def _clean(factory: async_sessionmaker[AsyncSession]) -> None:
    async with factory.begin() as session:
        for table in (
            PendingProductImageConfirmation,
            ProductImageRelation,
            ImageAssetOccurrence,
            ImageAsset,
            OutboxEvent,
            AuditEvent,
            KnowledgeIndexJob,
            DocumentParseRun,
            IngestionJob,
            IngestionSessionItem,
            IngestionSession,
            FileObject,
            User,
        ):
            await session.execute(delete(table))


async def _seed_file(
    factory: async_sessionmaker[AsyncSession],
    storage: LocalFileStorage,
    *,
    filename: str,
    data: bytes,
    with_run: bool = True,
) -> tuple[int, int, Path, int | None]:
    """Store one committed upload and return its file id, job id, blob path and parse run id."""

    async with factory.begin() as session:
        user = (
            await session.execute(select(User).where(User.telegram_user_id == TELEGRAM_USER_ID))
        ).scalar_one_or_none()
        if user is None:
            user = User(
                telegram_user_id=TELEGRAM_USER_ID, role=UserRole.USER, status=UserStatus.ACTIVE
            )
            session.add(user)
            await session.flush()
        ingestion_session = IngestionSession(
            user_id=user.id,
            status=IngestionSessionStatus.COMMITTED,
            expires_at=datetime.now(UTC),
        )
        session.add(ingestion_session)
        await session.flush()
        staged_path, staged_key = storage.create_staging_path(Path(filename).suffix.lower())
        staged_path.write_bytes(data)
        file_object = FileObject(
            storage_key=storage.activate(staged_key),
            original_filename=filename,
            content_type="application/octet-stream",
            byte_size=len(data),
            sha256=_digest(data),
            state=FileObjectState.ACTIVE,
        )
        session.add(file_object)
        job = IngestionJob(
            session_id=ingestion_session.id,
            status=IngestionJobStatus.PENDING,
            attempt_count=0,
        )
        session.add(job)
        await session.flush()
        session.add(
            IngestionSessionItem(
                session_id=ingestion_session.id,
                item_type=IngestionItemType.FILE,
                file_object_id=file_object.id,
                submitted_filename=filename,
                sender_user_id=user.id,
                telegram_chat_id=TELEGRAM_USER_ID,
                telegram_message_id=file_object.id,
                sequence_no=1,
            )
        )
        run_id: int | None = None
        if with_run:
            run = DocumentParseRun(
                file_object_id=file_object.id,
                ingestion_job_id=job.id,
                status=DocumentProcessingStatus.SUCCEEDED,
                parser_name="test-seed",
                parser_version="1.0",
            )
            session.add(run)
            await session.flush()
            run_id = run.id
        blob_path = storage.path_for_key(file_object.storage_key)
        return file_object.id, job.id, blob_path, run_id


async def _counts(factory: async_sessionmaker[AsyncSession]) -> tuple[int, int]:
    async with factory() as session:
        assets = (await session.execute(select(func.count(ImageAsset.id)))).scalar_one()
        occurrences = (
            await session.execute(select(func.count(ImageAssetOccurrence.id)))
        ).scalar_one()
        return int(assets), int(occurrences)


def _scenario(
    behavior: Callable[[async_sessionmaker[AsyncSession], LocalFileStorage], Awaitable[None]],
) -> None:
    async def run() -> None:
        settings = _settings_or_skip()
        root = Path(settings.file_storage_path) / "phase12a-test"
        storage = LocalFileStorage(root)
        engine = create_database_engine(settings)
        factory = create_session_factory(engine)
        try:
            await _clean(factory)
            await behavior(factory, storage)
        finally:
            await _clean(factory)
            await engine.dispose()
            shutil.rmtree(root, ignore_errors=True)

    asyncio.run(run())


def test_worker_registers_an_uploaded_photo_without_losing_the_original() -> None:
    async def behavior(
        factory: async_sessionmaker[AsyncSession], storage: LocalFileStorage
    ) -> None:
        # No seeded parse run: the worker must create the one the occurrence points at.
        file_id, job_id, blob_path, _ = await _seed_file(
            factory, storage, filename="telegram_photo.jpg", data=JPEG_BYTES, with_run=False
        )
        service = IngestionJobService(factory, storage)
        claim = await service.claim_next("phase12a-worker")
        assert claim is not None
        await service.process(claim)

        async with factory() as session:
            asset = (await session.execute(select(ImageAsset))).scalar_one()
            occurrence = (
                await session.execute(
                    select(ImageAssetOccurrence).where(
                        ImageAssetOccurrence.source_file_id == file_id
                    )
                )
            ).scalar_one()
            run = (
                await session.execute(
                    select(DocumentParseRun).where(DocumentParseRun.ingestion_job_id == job_id)
                )
            ).scalar_one()
            assert asset.content_sha256 == _digest(JPEG_BYTES)
            # A JPEG header with no frame marker states no size, so the row admits that instead of
            # guessing a thumbnail dimension.
            assert asset.format == "JPEG"
            assert (asset.width, asset.height) == (None, None)
            assert asset.storage_key.startswith("image_assets/")
            assert occurrence.image_asset_id == asset.id
            assert occurrence.origin_type == ImageAssetOrigin.STANDALONE_UPLOAD
            assert occurrence.source_parse_run_id == run.id
            assert occurrence.locator == {}
        assert blob_path.read_bytes() == JPEG_BYTES, "the uploaded original must survive"
        assert storage.path_for_key(asset.storage_key).read_bytes() == JPEG_BYTES

    _scenario(behavior)


def test_a_picture_inside_a_document_becomes_content_without_becoming_an_upload() -> None:
    async def behavior(
        factory: async_sessionmaker[AsyncSession], storage: LocalFileStorage
    ) -> None:
        logo = _png(120, 60)
        letterhead = _docx_bytes(logo)
        # No seeded parse run: the worker must create the one the occurrences point at.
        file_id, job_id, blob_path, _ = await _seed_file(
            factory,
            storage,
            filename="qiymet_siyahisi.docx",
            data=letterhead,
            with_run=False,
        )
        service = IngestionJobService(factory, storage)
        claim = await service.claim_next("phase12a-worker")
        assert claim is not None
        await service.process(claim)

        assert await _counts(factory) == (1, 2), "one picture, two places that show it"
        async with factory() as session:
            asset = (await session.execute(select(ImageAsset))).scalar_one()
            sighted = (await session.execute(select(ImageAssetOccurrence))).scalars().all()
            run = (
                await session.execute(
                    select(DocumentParseRun).where(DocumentParseRun.ingestion_job_id == job_id)
                )
            ).scalar_one()

            assert asset.content_sha256 == _digest(logo)
            assert (asset.mime_type, asset.width, asset.height) == ("image/png", 120, 60)
            assert asset.storage_key.startswith("image_assets/")
            assert run.status == DocumentProcessingStatus.SUCCEEDED
            # Text was readable, so a picture the build already describes must not defer the run.
            assert (run.requires_ocr, run.warning_count) == (False, 0)
            assert run.normalized_storage_key is not None
            described = storage.path_for_key(run.normalized_storage_key).read_bytes()

            for occurrence in sighted:
                assert occurrence.image_asset_id == asset.id
                assert occurrence.origin_type == ImageAssetOrigin.DOCX_EMBEDDED
                assert occurrence.source_file_id == file_id
                assert occurrence.source_parse_run_id == run.id
            paragraph = next(row for row in sighted if "paragraph_index" in row.locator)
            in_table = next(row for row in sighted if "table_index" in row.locator)
            assert paragraph.locator == {
                "paragraph_index": 0,
                "image_index": 1,
                "relationship_id": paragraph.locator["relationship_id"],
            }
            assert in_table.locator == {
                "table_index": 0,
                "row": 1,
                "column": 2,
                "image_index": 1,
                "relationship_id": paragraph.locator["relationship_id"],
            }

        # §23.3: the document is the only upload, and the picture is the only copy of its bytes.
        async with factory() as session:
            assert (await session.execute(select(func.count(FileObject.id)))).scalar_one() == 1
        assert blob_path.read_bytes() == letterhead, "the document keeps its own bytes"
        assert [path.name for path in storage.blob_root.rglob("*") if path.is_file()] == [
            blob_path.name
        ]
        copies = [
            path for path in storage.root.rglob("*") if path.is_file() and path.read_bytes() == logo
        ]
        assert [path.relative_to(storage.root).as_posix() for path in copies] == [asset.storage_key]
        assert PNG_START not in described, (
            "the parse record describes pictures, it never holds them"
        )

    _scenario(behavior)


def test_reprocessing_the_same_sighting_records_no_second_row() -> None:
    async def behavior(
        factory: async_sessionmaker[AsyncSession], storage: LocalFileStorage
    ) -> None:
        file_id, _, blob_path, run_id = await _seed_file(
            factory, storage, filename="sight.png", data=_png()
        )
        assert run_id is not None
        locator = {"page": 3, "image_index": 1}
        async with factory.begin() as session:
            registry = ImageAssetRegistry(session, storage)
            first = await registry.register(
                blob_path,
                origin=ImageAssetOrigin.PDF_EMBEDDED,
                source_file_id=file_id,
                source_parse_run_id=run_id,
                locator=locator,
            )
            assert first is not None
            assert (first.reused_content, first.recorded_occurrence) == (False, True)
            second = await registry.register(
                blob_path,
                origin=ImageAssetOrigin.PDF_EMBEDDED,
                source_file_id=file_id,
                source_parse_run_id=run_id,
                locator=locator,
            )
            assert second is not None
            assert (second.asset_id, second.reused_content, second.recorded_occurrence) == (
                first.asset_id,
                True,
                False,
            )

        assert await _counts(factory) == (1, 1)

    _scenario(behavior)


def test_one_logo_in_fifty_places_is_fifty_locators_and_one_copy_of_bytes() -> None:
    async def behavior(
        factory: async_sessionmaker[AsyncSession], storage: LocalFileStorage
    ) -> None:
        data = _png(120, 60)
        file_id, _, blob_path, run_id = await _seed_file(
            factory, storage, filename="letterhead.png", data=data
        )
        assert run_id is not None
        async with factory.begin() as session:
            registry = ImageAssetRegistry(session, storage)
            for page in range(1, 51):
                registration = await registry.register(
                    blob_path,
                    origin=ImageAssetOrigin.DOCX_EMBEDDED,
                    source_file_id=file_id,
                    source_parse_run_id=run_id,
                    locator={"table": 2, "cell": f"B{page}", "image_index": 0},
                )
                assert registration is not None

        assert await _counts(factory) == (1, 50)
        async with factory() as session:
            asset = (await session.execute(select(ImageAsset))).scalar_one()
            assert (asset.mime_type, asset.width, asset.height) == ("image/png", 120, 60)
        copies = [path for path in storage.image_asset_root.rglob("*") if path.is_file()]
        assert len(copies) == 1
        assert copies[0].read_bytes() == data

    _scenario(behavior)


def test_one_changed_image_byte_is_new_content_not_an_exact_duplicate() -> None:
    """Appearance does not decide byte identity or overwrite a distinct asset."""

    async def behavior(
        factory: async_sessionmaker[AsyncSession], storage: LocalFileStorage
    ) -> None:
        original = _png(120, 60)
        changed = original[:-1] + bytes([original[-1] ^ 1])
        assert len(original) == len(changed)
        async with factory.begin() as session:
            registry = ImageAssetRegistry(session, storage)
            first_id = await registry.ensure_content(original)
            duplicate_id = await registry.ensure_content(original)
            changed_id = await registry.ensure_content(changed)
            assert first_id is not None
            assert duplicate_id == first_id
            assert changed_id is not None and changed_id != first_id

        assert await _counts(factory) == (2, 0)
        stored = [path for path in storage.image_asset_root.rglob("*") if path.is_file()]
        assert len(stored) == 2
        assert {path.read_bytes() for path in stored} == {original, changed}

    _scenario(behavior)


def test_shared_bytes_outlive_a_document_while_its_provenance_cascades() -> None:
    async def behavior(
        factory: async_sessionmaker[AsyncSession], storage: LocalFileStorage
    ) -> None:
        # Two different documents that each embed the same logo: the document bytes differ, the
        # image content does not, which is exactly the case a per-file blob layout cannot express.
        logo = _png(80, 40)
        first_file, _, _, first_run = await _seed_file(
            factory, storage, filename="cover.pdf", data=b"%PDF-1.7 cover sheet"
        )
        second_file, _, _, second_run = await _seed_file(
            factory, storage, filename="appendix.xlsx", data=b"PK\x03\x04 appendix"
        )
        assert first_run is not None and second_run is not None
        sightings = (
            (
                first_file,
                first_run,
                _derived_image(storage, first_file, first_run, logo),
                ImageAssetOrigin.PDF_EMBEDDED,
            ),
            (
                second_file,
                second_run,
                _derived_image(storage, second_file, second_run, logo),
                ImageAssetOrigin.XLSX_EMBEDDED,
            ),
        )
        async with factory.begin() as session:
            registry = ImageAssetRegistry(session, storage)
            for position, (file_id, run_id, image_path, origin) in enumerate(sightings):
                registration = await registry.register(
                    image_path,
                    origin=origin,
                    source_file_id=file_id,
                    source_parse_run_id=run_id,
                    locator={"page": 1, "image_index": 0},
                )
                assert registration is not None
                assert registration.reused_content is (position > 0)

        async with factory.begin() as session:
            # Removing a source document means removing the rows that hang off it; the shared bytes
            # are not one document's property, so they stay for whoever else cites them.
            await session.execute(
                delete(DocumentParseRun).where(DocumentParseRun.file_object_id == first_file)
            )
            await session.execute(
                delete(IngestionSessionItem).where(
                    IngestionSessionItem.file_object_id == first_file
                )
            )
            await session.execute(delete(FileObject).where(FileObject.id == first_file))

        async with factory() as session:
            asset = (await session.execute(select(ImageAsset))).scalar_one()
            remaining = (await session.execute(select(ImageAssetOccurrence))).scalars().all()
            assert asset.byte_size == len(logo)
            assert [row.source_file_id for row in remaining] == [second_file]
        assert storage.path_for_key(asset.storage_key).is_file()

    _scenario(behavior)


def test_bytes_that_are_not_an_image_register_nothing_and_copy_nothing() -> None:
    async def behavior(
        factory: async_sessionmaker[AsyncSession], storage: LocalFileStorage
    ) -> None:
        data = b"a memo that no header claims to be an image"
        file_id, _, blob_path, run_id = await _seed_file(
            factory, storage, filename="memo.txt", data=data
        )
        assert run_id is not None
        async with factory.begin() as session:
            registry = ImageAssetRegistry(session, storage)
            registration = await registry.register(
                blob_path,
                origin=ImageAssetOrigin.STANDALONE_UPLOAD,
                source_file_id=file_id,
                source_parse_run_id=run_id,
            )

        assert registration is None
        assert await _counts(factory) == (0, 0)
        assert [path for path in storage.image_asset_root.rglob("*") if path.is_file()] == []

    _scenario(behavior)


def test_worker_prewarms_docx_embedded_images_after_committing_source() -> None:
    """An approved ingestion can precompute image vectors without importing Torch."""

    class StubEncoder:
        calls = 0

        def embed(self, paths: list[Path]) -> list[list[float]]:
            self.calls += len(paths)
            return [[1.0] + [0.0] * (CLIP_DIMENSION - 1) for _ in paths]

    async def behavior(
        factory: async_sessionmaker[AsyncSession], storage: LocalFileStorage
    ) -> None:
        image = Image.new("RGB", (128, 128), (200, 30, 20))
        ImageDraw.Draw(image).rectangle((20, 25, 102, 105), fill="blue")
        payload = BytesIO()
        image.save(payload, format="PNG")
        content = payload.getvalue()
        document = _docx_bytes(content)
        file_id, _, _, _ = await _seed_file(
            factory, storage, filename="embedded-photo.docx", data=document, with_run=False
        )
        cache = storage.root / "models-for-test"
        encoder = StubEncoder()
        semantic = SemanticVisualSearch(
            storage_root=storage.root,
            model_cache_root=cache,
            encoder=encoder,
        )
        service = IngestionJobService(factory, storage, visual_semantic=semantic)
        claim = await service.claim_next("docx-visual-test")
        assert claim is not None
        await service.process(claim)
        assert encoder.calls >= 1
        async with factory() as session:
            occurrences = (
                (
                    await session.execute(
                        select(ImageAssetOccurrence).where(
                            ImageAssetOccurrence.source_file_id == file_id
                        )
                    )
                )
                .scalars()
                .all()
            )
            assert len(occurrences) == 2
            asset = await session.get(ImageAsset, occurrences[0].image_asset_id)
            assert asset is not None
        digest = sha256(content).hexdigest()
        vector_root = cache / "visual-embeddings" / semantic.version / digest[:2]
        vector_file = vector_root / (digest + ".json")
        assert vector_file.is_file()
        assert digest in vector_file.read_text("utf-8")
        async with factory() as session:
            sources, truncated = await active_document_images(session)
            assert not truncated
            assert {source.occurrence_id for source in sources if source.source_file_id == file_id}
            assert all(
                source.origin == ImageAssetOrigin.DOCX_EMBEDDED
                for source in sources
                if source.source_file_id == file_id
            )
        delivery = SourceFileDeliveryService(factory, storage, max_bytes=25 * 1024 * 1024)
        original = await delivery.resolve_catalog_original(
            TELEGRAM_USER_ID, file_id=file_id, occurrence_id=occurrences[0].id
        )
        assert original.filename == "embedded-photo.docx"
        assert original.path.read_bytes() == document
        # A deleted original cannot stay visible merely because its vectors exist.
        async with factory.begin() as session:
            original = await session.get(FileObject, file_id)
            assert original is not None
            original.state = FileObjectState.DELETED
        async with factory() as session:
            sources, _ = await active_document_images(session)
            assert all(source.source_file_id != file_id for source in sources)
        with pytest.raises(FileDeliveryError):
            await delivery.resolve_catalog_original(
                TELEGRAM_USER_ID, file_id=file_id, occurrence_id=occurrences[0].id
            )

    _scenario(behavior)


@pytest.mark.parametrize("failure_mode", ["disabled", "index_unavailable"])
def test_committed_photo_invalidates_ready_visual_index_on_failed_ingestion(
    failure_mode: str,
) -> None:
    """A durable new photo must never leave a stale full-corpus READY marker."""

    class FailingIndex:
        async def ensure(self) -> None:
            raise RuntimeError("synthetic Qdrant outage")

    async def behavior(
        factory: async_sessionmaker[AsyncSession], storage: LocalFileStorage
    ) -> None:
        marker_key = "office_visual_dino_ann_ready_v1"
        async with factory.begin() as session:
            previous = await session.get(SystemMetadata, marker_key)
            old_value = previous.value if previous is not None else None
            if previous is None:
                session.add(SystemMetadata(key=marker_key, value="ready"))
            else:
                previous.value = "ready"
        try:
            file_id, job_id, _, _ = await _seed_file(
                factory,
                storage,
                filename="test-ready-invalidated.jpg",
                data=JPEG_BYTES,
                with_run=False,
            )
            semantic = (
                SemanticVisualSearch(
                    storage_root=storage.root,
                    model_cache_root=storage.root / "test-model-cache",
                )
                if failure_mode == "index_unavailable"
                else None
            )
            service = IngestionJobService(
                factory,
                storage,
                visual_semantic=semantic,
                visual_ann=(
                    cast(VisualANNIndex, FailingIndex())
                    if failure_mode == "index_unavailable"
                    else None
                ),
            )
            claim = await service.claim_next("visual-ready-failure")
            assert claim is not None
            await service.process(claim)
            async with factory() as session:
                original = await session.get(FileObject, file_id)
                assert original is not None
                assert original.state == FileObjectState.ACTIVE
                recorded = await session.get(SystemMetadata, marker_key)
                assert recorded is not None and recorded.value == "stale"
                job = await session.get(IngestionJob, job_id)
                assert job is not None
                assert job.status in {IngestionJobStatus.SUCCEEDED, IngestionJobStatus.PARTIAL}
                assert (
                    await session.scalar(
                        select(ImageAssetOccurrence.id)
                        .where(ImageAssetOccurrence.source_file_id == file_id)
                        .limit(1)
                    )
                    is not None
                )
        finally:
            async with factory.begin() as session:
                prior = await session.get(SystemMetadata, marker_key)
                if prior is not None:
                    if old_value is None:
                        await session.delete(prior)
                    else:
                        prior.value = old_value

    _scenario(behavior)
